import sqlite3
from datetime import datetime, timedelta
import logging

# 导入日志配置
try:
    from logging_config import logger, log_database_operation
except ImportError:
    # 如果 logging_config 不存在，创建简单的logger
    logging.basicConfig(level=logging.INFO)
    logger = logging.getLogger('PatientReminderSystem')
    def log_database_operation(*args, **kwargs):
        pass

DB_NAME = "patients.db"

def init_db():
    """初始化数据库，创建患者表、回访计划表和提醒配置表"""
    logger.info("数据库初始化开始")
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    
    try:
        # 创建患者表
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS patients (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                gender TEXT,
                hospital_number TEXT,
                discharge_date DATE,
                diagnosis TEXT,
                contact_person TEXT,
                contact_phone TEXT,
                status TEXT DEFAULT 'pending',
                created_at TIMESTAMP DEFAULT (datetime('now', 'localtime'))
            )
        ''')
        
        # 创建回访计划表（支持多次回访）
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS visit_plans (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                patient_id INTEGER NOT NULL,
                visit_date DATE NOT NULL,
                visit_type TEXT DEFAULT '常规回访',
                status TEXT DEFAULT 'pending',
                notes TEXT,
                completed_at TIMESTAMP,
                created_at TIMESTAMP DEFAULT (datetime('now', 'localtime')),
                FOREIGN KEY (patient_id) REFERENCES patients(id) ON DELETE CASCADE
            )
        ''')
        
        # 创建回访记录表（存储详细的回访内容）
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS visit_records (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                plan_id INTEGER NOT NULL,
                patient_id INTEGER NOT NULL,
                visit_date DATE NOT NULL,
                record_content TEXT,
                patient_feedback TEXT,
                health_status TEXT,
                medication_info TEXT,
                next_visit_date DATE,
                recorder TEXT,
                created_at TIMESTAMP DEFAULT (datetime('now', 'localtime')),
                FOREIGN KEY (plan_id) REFERENCES visit_plans(id) ON DELETE CASCADE,
                FOREIGN KEY (patient_id) REFERENCES patients(id) ON DELETE CASCADE
            )
        ''')
        
        # 创建提醒配置表
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS reminder_settings (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                setting_key TEXT UNIQUE NOT NULL,
                setting_value TEXT NOT NULL,
                description TEXT,
                updated_at TIMESTAMP DEFAULT (datetime('now', 'localtime'))
            )
        ''')
        
        # 插入默认提醒配置
        default_settings = [
            ('advance_days', '0', '提前N天提醒（0表示当天提醒）'),
            ('reminder_time', '09:00', '每天固定提醒时间（24小时制）'),
            ('repeat_interval', '24', '重复提醒间隔（小时）'),
            ('enable_advance_reminder', 'false', '是否启用提前提醒'),
            ('enable_repeat_reminder', 'true', '是否启用重复提醒'),
        ]
        
        for key, value, desc in default_settings:
            cursor.execute('''
                INSERT OR IGNORE INTO reminder_settings (setting_key, setting_value, description)
                VALUES (?, ?, ?)
            ''', (key, value, desc))
        
        conn.commit()
        logger.info("数据库初始化成功")
        log_database_operation('INIT', 'all tables', success=True)
        
    except Exception as e:
        logger.error(f"数据库初始化失败: {str(e)}", exc_info=True)
        log_database_operation('INIT', 'all tables', success=False, error_msg=str(e))
        raise
    finally:
        conn.close()

def add_patient(name, gender, discharge_date, visit_dates, diagnosis, contact_person, contact_phone, hospital_number=None):
    """
    添加新患者
    visit_dates: 回访日期列表，例如 ['2026-05-10', '2026-06-10', '2026-08-10']
    hospital_number: 住院号（可选）
    """
    logger.info(f"添加患者 | 姓名: {name} | 住院号: {hospital_number or '无'}")
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    
    try:
        # 使用本地时间
        local_time = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        
        # 插入患者信息
        cursor.execute('''
            INSERT INTO patients (name, gender, hospital_number, discharge_date, diagnosis, contact_person, contact_phone, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        ''', (name, gender, hospital_number, discharge_date, diagnosis, contact_person, contact_phone, local_time))
        
        # 获取患者ID
        patient_id = cursor.lastrowid
        
        # 插入多个回访计划
        for visit_date in visit_dates:
            cursor.execute('''
                INSERT INTO visit_plans (patient_id, visit_date, status)
                VALUES (?, ?, 'pending')
            ''', (patient_id, visit_date))
        
        conn.commit()
        logger.info(f"患者添加成功 | ID: {patient_id} | 回访计划数: {len(visit_dates)}")
        log_database_operation('INSERT', 'patients + visit_plans', success=True)
        
    except Exception as e:
        logger.error(f"患者添加失败 | 姓名: {name} | 错误: {str(e)}", exc_info=True)
        log_database_operation('INSERT', 'patients + visit_plans', success=False, error_msg=str(e))
        raise
    finally:
        conn.close()

def get_due_patients():
    """查询今天需要回访的患者（不包括明天）"""
    today = datetime.now().strftime('%Y-%m-%d')
    
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    
    # 从回访计划表中查询今天的任务
    cursor.execute('''
        SELECT p.name, p.gender, p.diagnosis, vp.visit_date, p.contact_person, p.contact_phone, vp.id as plan_id
        FROM visit_plans vp
        JOIN patients p ON vp.patient_id = p.id
        WHERE vp.visit_date = ? AND vp.status != 'completed'
        ORDER BY vp.visit_date ASC
    ''', (today,))
    
    plans = cursor.fetchall()
    conn.close()
    
    result = []
    for p in plans:
        result.append({
            "name": p[0],
            "gender": p[1],
            "diagnosis": p[2],
            "visit_date": p[3],
            "contact_person": p[4],
            "contact_phone": p[5],
            "plan_id": p[6],
            "urgency": "今日紧急"
        })
    return result

def get_all_patients():
    """获取所有患者数据及其回访计划"""
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    
    # 获取所有患者
    cursor.execute('''
        SELECT id, name, gender, hospital_number, discharge_date, diagnosis, contact_person, contact_phone, status, created_at
        FROM patients
        ORDER BY created_at DESC
    ''')
    
    patients = cursor.fetchall()
    
    result = []
    for p in patients:
        # 获取该患者的所有回访计划
        cursor.execute('''
            SELECT id, visit_date, visit_type, status, completed_at
            FROM visit_plans
            WHERE patient_id = ?
            ORDER BY visit_date ASC
        ''', (p[0],))
        
        visit_plans = cursor.fetchall()
        visit_dates = [vp[1] for vp in visit_plans]
        
        # 计算患者的整体状态（如果有任何一个计划未完成，则患者状态为 pending）
        has_pending = any(vp[3] == 'pending' for vp in visit_plans)
        all_completed = all(vp[3] == 'completed' for vp in visit_plans) if visit_plans else False
        
        if all_completed and visit_plans:
            overall_status = 'completed'
        elif has_pending:
            overall_status = 'pending'
        else:
            overall_status = p[8]  # 使用原状态
        
        result.append({
            "id": p[0],
            "name": p[1],
            "gender": p[2],
            "hospital_number": p[3],
            "discharge_date": p[4],
            "visit_dates": visit_dates,  # 多个回访日期
            "visit_date": visit_dates[0] if visit_dates else None,  # 第一个回访日期（兼容旧代码）
            "diagnosis": p[5],
            "contact_person": p[6],
            "contact_phone": p[7],
            "status": overall_status,  # 根据回访计划计算的总体状态
            "created_at": p[9],
            "visit_plans": [{
                "plan_id": vp[0],
                "visit_date": vp[1],
                "visit_type": vp[2],
                "status": vp[3],
                "completed_at": vp[4]
            } for vp in visit_plans]  # 完整的回访计划列表
        })
    
    conn.close()
    return result

def update_patient_status(patient_id, status):
    """更新患者所有回访计划的状态"""
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    
    # 更新该患者的所有回访计划
    cursor.execute('''
        UPDATE visit_plans SET status = ?, completed_at = CASE WHEN ? = 'completed' THEN datetime('now', 'localtime') ELSE NULL END
        WHERE patient_id = ?
    ''', (status, status, patient_id))
    
    conn.commit()
    conn.close()

def update_visit_plan_status(plan_id, status):
    """更新单个回访计划的状态"""
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    
    cursor.execute('''
        UPDATE visit_plans SET status = ?, completed_at = CASE WHEN ? = 'completed' THEN datetime('now', 'localtime') ELSE NULL END
        WHERE id = ?
    ''', (status, status, plan_id))
    
    conn.commit()
    conn.close()

def add_visit_plan(patient_id, visit_date, visit_type='常规回访'):
    """为患者添加新的回访计划"""
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    
    cursor.execute('''
        INSERT INTO visit_plans (patient_id, visit_date, visit_type, status)
        VALUES (?, ?, ?, 'pending')
    ''', (patient_id, visit_date, visit_type))
    
    conn.commit()
    conn.close()

def import_patients_from_csv(csv_data):
    """
    从CSV数据批量导入患者
    csv_data: list of dicts，每个dict包含患者信息
    返回: (success_count, error_count, errors_list)
    """
    logger.info(f"CSV批量导入开始 | 记录数: {len(csv_data)}")
    success_count = 0
    error_count = 0
    errors = []
    
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    
    try:
        for i, row in enumerate(csv_data, 1):
            try:
                # 提取字段（转换为字符串，处理数字类型）
                name = str(row.get('姓名', '')).strip()
                gender = str(row.get('性别', '')).strip()
                hospital_number = str(row.get('住院号', '')).strip() if row.get('住院号') else ''
                discharge_date = str(row.get('出院日期', '')).strip()
                visit_date = str(row.get('回访日期', '')).strip()
                diagnosis = str(row.get('诊断', '')).strip()
                contact_person = str(row.get('联系人', '')).strip()
                contact_phone = str(row.get('联系电话', '')).strip()
                status = str(row.get('状态', 'pending')).strip()
                
                # 验证必填字段
                if not all([name, gender, discharge_date, visit_date, diagnosis, contact_person, contact_phone]):
                    error_count += 1
                    errors.append(f"第{i}行：缺少必填字段")
                    continue
                
                # 插入患者信息
                local_time = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                cursor.execute('''
                    INSERT INTO patients (name, gender, hospital_number, discharge_date, diagnosis, contact_person, contact_phone, status, created_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''', (name, gender, hospital_number if hospital_number else None, discharge_date, diagnosis, contact_person, contact_phone, status, local_time))
                
                patient_id = cursor.lastrowid
                
                # 创建回访计划
                cursor.execute('''
                    INSERT INTO visit_plans (patient_id, visit_date, status)
                    VALUES (?, ?, ?)
                ''', (patient_id, visit_date, status))
                
                success_count += 1
                
            except Exception as e:
                error_count += 1
                error_msg = f"第{i}行 ({row.get('姓名', '未知')}): {str(e)}"
                errors.append(error_msg)
                logger.warning(f"CSV导入跳过记录 | {error_msg}")
        
        conn.commit()
        logger.info(f"CSV批量导入完成 | 成功: {success_count} | 失败: {error_count}")
        log_database_operation('BATCH_INSERT', 'patients + visit_plans (CSV)', success=True)
        
    except Exception as e:
        conn.rollback()
        logger.error(f"CSV批量导入失败 | 错误: {str(e)}", exc_info=True)
        log_database_operation('BATCH_INSERT', 'patients + visit_plans (CSV)', success=False, error_msg=str(e))
        raise
    finally:
        conn.close()
    
    return success_count, error_count, errors

def get_overdue_patients():
    """查询已过回访日期但状态仍为pending的患者（忘记回访）"""
    today = datetime.now().strftime('%Y-%m-%d')
    
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    cursor.execute('''
        SELECT p.id, p.name, p.gender, p.diagnosis, vp.visit_date, p.contact_person, p.contact_phone, p.created_at, vp.id as plan_id
        FROM visit_plans vp
        JOIN patients p ON vp.patient_id = p.id
        WHERE vp.visit_date < ? AND vp.status = 'pending'
        ORDER BY vp.visit_date ASC
    ''', (today,))
    
    plans = cursor.fetchall()
    conn.close()
    
    result = []
    for p in plans:
        # 计算 overdue 天数
        days_overdue = (datetime.now() - datetime.strptime(p[4], '%Y-%m-%d')).days
        result.append({
            "id": p[0],
            "name": p[1],
            "gender": p[2],
            "diagnosis": p[3],
            "visit_date": p[4],
            "contact_person": p[5],
            "contact_phone": p[6],
            "created_at": p[7],
            "plan_id": p[8],
            "days_overdue": days_overdue
        })
    return result

def find_duplicate_patients_by_hospital_number():
    """
    查找具有相同住院号的重复患者记录
    返回: dict，key为住院号，value为该住院号对应的患者列表
    """
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    
    # 查找有重复住院号的患者（排除空值）
    cursor.execute('''
        SELECT hospital_number, COUNT(*) as count
        FROM patients
        WHERE hospital_number IS NOT NULL AND hospital_number != ''
        GROUP BY hospital_number
        HAVING count > 1
    ''')
    
    duplicates = cursor.fetchall()
    
    result = {}
    for hosp_num, count in duplicates:
        cursor.execute('''
            SELECT id, name, gender, discharge_date, diagnosis, contact_person, contact_phone, status, created_at
            FROM patients
            WHERE hospital_number = ?
            ORDER BY created_at DESC
        ''', (hosp_num,))
        
        patients = cursor.fetchall()
        result[hosp_num] = [{
            "id": p[0],
            "name": p[1],
            "gender": p[2],
            "discharge_date": p[3],
            "diagnosis": p[4],
            "contact_person": p[5],
            "contact_phone": p[6],
            "status": p[7],
            "created_at": p[8]
        } for p in patients]
    
    conn.close()
    return result

def merge_duplicate_patients(hospital_number, keep_patient_id):
    """
    合并具有相同住院号的重复患者记录
    
    参数:
        hospital_number: 住院号
        keep_patient_id: 要保留的患者ID（主记录）
    
    返回:
        (merged_count, deleted_count, message)
    """
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    
    try:
        # 获取该住院号下的所有患者
        cursor.execute('''
            SELECT id FROM patients WHERE hospital_number = ?
        ''', (hospital_number,))
        
        all_patient_ids = [row[0] for row in cursor.fetchall()]
        
        if keep_patient_id not in all_patient_ids:
            return (0, 0, f"错误：患者ID {keep_patient_id} 不属于住院号 {hospital_number}")
        
        # 找出需要合并的患者（除了保留的那个）
        merge_patient_ids = [pid for pid in all_patient_ids if pid != keep_patient_id]
        
        if not merge_patient_ids:
            return (0, 0, "没有需要合并的重复记录")
        
        merged_count = 0
        deleted_count = 0
        
        # 对于每个需要合并的患者
        for old_patient_id in merge_patient_ids:
            # 1. 迁移回访计划到保留的患者
            cursor.execute('''
                UPDATE visit_plans SET patient_id = ? WHERE patient_id = ?
            ''', (keep_patient_id, old_patient_id))
            
            migrated_plans = cursor.rowcount
            merged_count += migrated_plans
            
            # 2. 删除旧的患者记录
            cursor.execute('DELETE FROM patients WHERE id = ?', (old_patient_id,))
            deleted_count += 1
        
        conn.commit()
        
        message = f"成功合并住院号 {hospital_number} 的重复记录：\n"
        message += f"- 保留了患者ID: {keep_patient_id}\n"
        message += f"- 合并了 {merged_count} 个回访计划\n"
        message += f"- 删除了 {deleted_count} 条重复患者记录"
        
        return (merged_count, deleted_count, message)
        
    except Exception as e:
        conn.rollback()
        return (0, 0, f"合并失败: {str(e)}")
    finally:
        conn.close()

def add_visit_record(plan_id, patient_id, visit_date, record_content='', patient_feedback='', 
                    health_status='', medication_info='', next_visit_date=None, recorder=''):
    """
    添加回访记录
    
    参数:
        plan_id: 回访计划ID
        patient_id: 患者ID
        visit_date: 回访日期
        record_content: 回访内容
        patient_feedback: 患者反馈
        health_status: 健康状况
        medication_info: 用药情况
        next_visit_date: 下次回访日期
        recorder: 记录人
    """
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    
    local_time = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    
    cursor.execute('''
        INSERT INTO visit_records (plan_id, patient_id, visit_date, record_content, 
                                  patient_feedback, health_status, medication_info, 
                                  next_visit_date, recorder, created_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (plan_id, patient_id, visit_date, record_content, patient_feedback, 
          health_status, medication_info, next_visit_date, recorder, local_time))
    
    conn.commit()
    conn.close()

def get_visit_records(patient_id, plan_id=None):
    """
    获取患者的回访记录
    
    参数:
        patient_id: 患者ID
        plan_id: 回访计划ID（可选，如果提供则只返回该计划的记录）
    
    返回:
        回访记录列表
    """
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    
    if plan_id:
        cursor.execute('''
            SELECT id, plan_id, patient_id, visit_date, record_content, 
                   patient_feedback, health_status, medication_info, 
                   next_visit_date, recorder, created_at
            FROM visit_records
            WHERE patient_id = ? AND plan_id = ?
            ORDER BY visit_date DESC
        ''', (patient_id, plan_id))
    else:
        cursor.execute('''
            SELECT id, plan_id, patient_id, visit_date, record_content, 
                   patient_feedback, health_status, medication_info, 
                   next_visit_date, recorder, created_at
            FROM visit_records
            WHERE patient_id = ?
            ORDER BY visit_date DESC
        ''', (patient_id,))
    
    records = cursor.fetchall()
    conn.close()
    
    return [{
        "id": r[0],
        "plan_id": r[1],
        "patient_id": r[2],
        "visit_date": r[3],
        "record_content": r[4],
        "patient_feedback": r[5],
        "health_status": r[6],
        "medication_info": r[7],
        "next_visit_date": r[8],
        "recorder": r[9],
        "created_at": r[10]
    } for r in records]

def get_patient_full_profile(patient_id):
    """
    获取患者的完整档案信息（包括所有回访记录和计划）
    
    参数:
        patient_id: 患者ID
    
    返回:
        包含患者基本信息、回访计划、回访记录的字典
    """
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    
    # 获取患者基本信息
    cursor.execute('''
        SELECT id, name, gender, hospital_number, discharge_date, diagnosis, 
               contact_person, contact_phone, status, created_at
        FROM patients
        WHERE id = ?
    ''', (patient_id,))
    
    patient = cursor.fetchone()
    
    if not patient:
        conn.close()
        return None
    
    # 获取回访计划
    cursor.execute('''
        SELECT id, visit_date, visit_type, status, notes, completed_at, created_at
        FROM visit_plans
        WHERE patient_id = ?
        ORDER BY visit_date ASC
    ''', (patient_id,))
    
    plans = cursor.fetchall()
    
    # 获取回访记录
    cursor.execute('''
        SELECT id, plan_id, visit_date, record_content, patient_feedback, 
               health_status, medication_info, next_visit_date, recorder, created_at
        FROM visit_records
        WHERE patient_id = ?
        ORDER BY visit_date DESC
    ''', (patient_id,))
    
    records = cursor.fetchall()
    conn.close()
    
    return {
        "patient": {
            "id": patient[0],
            "name": patient[1],
            "gender": patient[2],
            "hospital_number": patient[3],
            "discharge_date": patient[4],
            "diagnosis": patient[5],
            "contact_person": patient[6],
            "contact_phone": patient[7],
            "status": patient[8],
            "created_at": patient[9]
        },
        "visit_plans": [{
            "plan_id": p[0],
            "visit_date": p[1],
            "visit_type": p[2],
            "status": p[3],
            "notes": p[4],
            "completed_at": p[5],
            "created_at": p[6]
        } for p in plans],
        "visit_records": [{
            "record_id": r[0],
            "plan_id": r[1],
            "visit_date": r[2],
            "record_content": r[3],
            "patient_feedback": r[4],
            "health_status": r[5],
            "medication_info": r[6],
            "next_visit_date": r[7],
            "recorder": r[8],
            "created_at": r[9]
        } for r in records]
    }

def export_patient_to_excel(patient_id, output_path):
    """
    导出患者档案到Excel文件
    
    参数:
        patient_id: 患者ID
        output_path: 输出文件路径
    
    返回:
        (success, message)
    """
    try:
        import pandas as pd
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
        
        # 获取患者完整档案
        profile = get_patient_full_profile(patient_id)
        
        if not profile:
            return (False, "患者不存在")
        
        # 创建工作簿
        wb = Workbook()
        
        # 1. 患者基本信息表
        ws1 = wb.active
        ws1.title = "患者基本信息"
        
        # 设置标题样式
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        
        # 写入基本信息
        patient = profile['patient']
        basic_info = [
            ['字段', '值'],
            ['患者ID', patient['id']],
            ['姓名', patient['name']],
            ['性别', patient['gender']],
            ['住院号', patient['hospital_number'] or '无'],
            ['出院日期', patient['discharge_date']],
            ['诊断', patient['diagnosis']],
            ['联系人', patient['contact_person']],
            ['联系电话', patient['contact_phone']],
            ['状态', '已完成' if patient['status'] == 'completed' else '待回访'],
            ['建档时间', patient['created_at']]
        ]
        
        for row_idx, row_data in enumerate(basic_info, 1):
            for col_idx, value in enumerate(row_data, 1):
                cell = ws1.cell(row=row_idx, column=col_idx, value=value)
                if row_idx == 1:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = header_alignment
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
        
        # 调整列宽
        ws1.column_dimensions['A'].width = 15
        ws1.column_dimensions['B'].width = 30
        
        # 2. 回访计划表
        ws2 = wb.create_sheet("回访计划")
        
        plans_header = ['计划ID', '回访日期', '回访类型', '状态', '备注', '完成时间', '创建时间']
        ws2.append(plans_header)
        
        # 设置表头样式
        for cell in ws2[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        
        # 写入回访计划数据
        for plan in profile['visit_plans']:
            ws2.append([
                plan['plan_id'],
                plan['visit_date'],
                plan['visit_type'],
                '已完成' if plan['status'] == 'completed' else '待回访',
                plan['notes'] or '',
                plan['completed_at'] or '',
                plan['created_at']
            ])
        
        # 调整列宽
        for col in ws2.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            ws2.column_dimensions[column].width = adjusted_width
        
        # 3. 回访记录表
        ws3 = wb.create_sheet("回访记录")
        
        records_header = ['记录ID', '回访日期', '回访内容', '患者反馈', '健康状况', '用药情况', '下次回访日期', '记录人', '记录时间']
        ws3.append(records_header)
        
        # 设置表头样式
        for cell in ws3[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        
        # 写入回访记录数据
        for record in profile['visit_records']:
            ws3.append([
                record['record_id'],
                record['visit_date'],
                record['record_content'] or '',
                record['patient_feedback'] or '',
                record['health_status'] or '',
                record['medication_info'] or '',
                record['next_visit_date'] or '',
                record['recorder'] or '',
                record['created_at']
            ])
        
        # 调整列宽
        for col in ws3.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 40)
            ws3.column_dimensions[column].width = adjusted_width
        
        # 保存文件
        wb.save(output_path)
        
        return (True, f"成功导出到: {output_path}")
        
    except ImportError:
        return (False, "需要安装 openpyxl 库: pip install openpyxl")
    except Exception as e:
        return (False, f"导出失败: {str(e)}")

def get_reminder_settings():
    """
    获取所有提醒配置
    
    返回:
        dict: 提醒配置字典
    """
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    
    cursor.execute('SELECT setting_key, setting_value, description FROM reminder_settings')
    settings = cursor.fetchall()
    conn.close()
    
    result = {}
    for key, value, desc in settings:
        result[key] = {
            'value': value,
            'description': desc
        }
    
    return result

def update_reminder_setting(setting_key, setting_value):
    """
    更新单个提醒配置
    
    参数:
        setting_key: 配置键
        setting_value: 配置值
    
    返回:
        bool: 是否成功
    """
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    
    local_time = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    
    cursor.execute('''
        UPDATE reminder_settings 
        SET setting_value = ?, updated_at = ?
        WHERE setting_key = ?
    ''', (setting_value, local_time, setting_key))
    
    success = cursor.rowcount > 0
    conn.commit()
    conn.close()
    
    return success

def update_reminder_settings_batch(settings_dict):
    """
    批量更新提醒配置
    
    参数:
        settings_dict: 配置字典 {key: value}
    
    返回:
        (success_count, error_count)
    """
    success_count = 0
    error_count = 0
    
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    
    local_time = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    
    for key, value in settings_dict.items():
        try:
            cursor.execute('''
                UPDATE reminder_settings 
                SET setting_value = ?, updated_at = ?
                WHERE setting_key = ?
            ''', (str(value), local_time, key))
            
            if cursor.rowcount > 0:
                success_count += 1
            else:
                error_count += 1
        except Exception as e:
            error_count += 1
    
    conn.commit()
    conn.close()
    
    return success_count, error_count

# 初始化数据库
init_db()
